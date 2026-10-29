#!/usr/bin/env python3
"""从结果展示表生成 SSOT JSON 与自然语言描述（markdown + docx）。

用法：
  python3 utils/build_reports.py <结果展示表.xlsx> --workloads 9 -o reports/9load
"""

import argparse
import importlib.util
import json
import os
import re

import openpyxl
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
META_PATH = os.path.join(HERE, "board_meta.json")


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


scoring = _load("scoring", os.path.join(HERE, "score_seven_workload.py"))
excel2json = _load("excel2json", os.path.join(HERE, "excel_to_result_json.py"))

WORKLOAD_LABEL = {
    "pid": "PID伺服电机控制", "ekf": "EKF飞控姿态解算", "fast": "FAST角点检测",
    "icp": "ICP点云配准", "epnp": "EPnP位姿求解", "modbus": "Modbus工业协议通信",
    "mqtt": "MQTT消息通信", "ewma": "EWMA异常检测", "cusum": "CUSUM异常检测",
}
CATEGORY_LABEL = {
    "realtime_control": "实时控制", "perception": "感知计算",
    "communication": "通信协同", "monitoring": "运维监控",
}
CAT_IN_TEXT = [
    ("realtime_control", "在实时控制类负载中"),
    ("perception", "在感知计算类负载中"),
    ("communication", "在通信协同类负载中"),
    ("monitoring", "在运维监控类负载中"),
]
OP_LABEL = {
    "sem_take": "信号量获取", "sem_release": "信号量释放",
    "mq_send": "消息发送", "mq_recv": "消息接收",
    "mutex_take": "互斥锁获取", "mutex_release": "互斥锁释放",
    "mempool_alloc": "内存块申请", "mempool_free": "内存块释放",
}
SCENE_LABEL = [("immediate_us", "立即执行"), ("suspend_us", "挂起睡眠"),
               ("low_prio_us", "低优先级就绪"), ("high_prio_us", "高优先级恢复")]


def _fmt(v, digits=3):
    return "—" if v is None else f"{v:.{digits}f}"


def _pct(v):
    return "—" if v is None else f"{v * 100:.2f}%"


def build_narrative(data, score):
    env = data["env"]
    board = env["board"]
    rt = data["modules"]["test-realtime"]
    sc = data["modules"]["test-schedule"]
    wl = data["modules"]["typical-workload"]["workloads"]
    power = data["power"]["items"]

    blocks = []
    blocks.append(("h1", f"一、测试结果描述（{board}）"))
    blocks.append(("p", f"本次针对{board}开展的工业实时操作系统基准测试，主要涵盖基础能力、典型负载性能、实时性能和功耗四个方面。"))
    os_version = env.get("os_version") or ""
    bsp_version = env.get("bsp_version") or ""
    if os_version or bsp_version:
        detail = " ".join(x for x in [os_version, bsp_version] if x)
        blocks.append(("p", f"被测系统为 {env.get('os_name')} {detail}。"))

    blocks.append(("h2", "（1）基础能力测试"))
    blocks.append(("p", "基础能力主要考察操作系统的内存访问、文件系统读写、网络通信及POSIX接口兼容性，"
                        "包括内存带宽、内存延迟、小文件IOPS、大文件IOPS、网络带宽、网络延迟和接口兼容性共7项指标。"
                        "目前该部分暂无实测数据，评分按占位值60分计。"))

    blocks.append(("h2", "（2）典型负载性能测试"))
    parts = []
    for cat, prefix in CAT_IN_TEXT:
        items = [w for w in wl if w["category"] == cat]
        if not items:
            continue
        text = "；".join(f"{WORKLOAD_LABEL.get(w['name'], w['name'])}任务为 {w['exec_time_ms']:.3f} ms"
                        for w in items)
        parts.append(f"{prefix}{text}")
    blocks.append(("p", "典型负载涵盖实时控制、感知计算、通信协同和运维监控四类应用场景，"
                        "主要通过典型任务的执行时间评价系统对工业应用的支撑能力。"
                        + "；".join(parts) + "。上述指标反映系统在不同工业任务场景下的执行效率，执行时间越短越好。"))

    blocks.append(("h2", "（3）实时性能测试"))
    blocks.append(("p", "实时性能测试包括基础实时性能和可调度性能。其中："))
    single = rt["single_core"]
    blocks.append(("li", f"1) 系统上下文切换延迟为 {_fmt(single['context_switch']['avg_us'])} μs；"))
    blocks.append(("li", f"2) 中断响应延迟为 {_fmt(single['interrupt']['avg_us'])} μs；"))
    ops = single.get("service_cost") or []
    if ops:
        segs = []
        for op in ops:
            vals = [f"{label} {_fmt(op[key])} μs" for key, label in SCENE_LABEL if op.get(key) is not None]
            if vals:
                segs.append(f"{OP_LABEL.get(op['operation'], op['operation'])}（" + "、".join(vals) + "）")
        scene_count = sum(1 for op in ops for key, _ in SCENE_LABEL if op.get(key) is not None)
        blocks.append(("li", f"3) 任务调度延迟共测试{scene_count}个场景：" + "；".join(segs) + "；"))
    mc = rt.get("multicore") or {}
    access = mc.get("access") or {}
    if access:
        segs = [f"{op} {', '.join('—' if v is None else str(v) for v in vals)} GB/s"
                for op, vals in access.items()]
        blocks.append(("li", "4) 多核存取性能按1/2/4/8并发度依次为：" + "；".join(segs) + "。"))

    blocks.append(("h3", "可调度性能"))
    grads = sc["gradients"]
    rates = [g["miss_rate"] for g in grads if g["miss_rate"] is not None]
    if rates:
        detail = "、".join(f"{g['utilization_percent']}% 为 {_pct(g['miss_rate'])}" for g in grads)
        nonzero = [(g['utilization_percent'], g['miss_rate']) for g in grads
                   if g['miss_rate'] is not None and g['miss_rate'] > 0]
        if nonzero:
            hit = "、".join(f"{u}%" for u, _ in nonzero)
            peak = max(r for _, r in nonzero)
            verdict = f"其中 {hit} 梯度出现截止时间错失，最高错失率 {_pct(peak)}"
        else:
            verdict = "所有梯度均未出现截止时间错失"
        blocks.append(("p",
            f"可调度性能评价系统在不同处理器利用率下保障周期性任务截止时间的能力。"
            f"测试采用30%至100%、每10%一个梯度的任务集配置，共{len(grads)}组结果，"
            f"各梯度错失率依次为：{detail}。{verdict}。"
            f"平均截止时间错失率为 {_pct(sc['summary']['average_miss_rate'])}，"
            f"可调度性能得分为 {_fmt(sc['summary']['final_score'], 2)} 分。"))

    blocks.append(("h2", "（4）功耗测试"))
    pw = {item["item"]: item for item in power}
    blocks.append(("p",
        "功耗测试覆盖静态空闲、CPU压力、内存压力和文件系统压力四类工况。"
        f"静态空闲能耗 {_fmt(pw.get('static', {}).get('energy_j'), 3)} J，"
        f"CPU压力 {_fmt(pw.get('cpu', {}).get('energy_j'), 3)} J，"
        f"内存压力 {_fmt(pw.get('memory', {}).get('energy_j'), 3)} J，"
        f"文件系统压力 {_fmt(pw.get('file', {}).get('energy_j'), 3)} J。"))

    blocks.append(("h1", f"二、测试评分描述（{board}）"))
    blocks.append(("p", "基于工业操作系统基准测试评分体系，按照基础能力20%、典型负载性能40%、"
                        "实时性能30%和功耗10%的权重计算综合得分。"))
    blocks.append(("p", "评分公式：总分 = 0.2×F + 0.3×RT + 0.4×T + 0.1×P。"))
    cats = score["categories"]
    blocks.append(("li", f"1) 基础能力（权重20%，F）：{score['F']:.2f} 分。该分数为占位值，"
                         "后续需结合内存、文件系统、网络和POSIX接口兼容性测试结果确定。"))
    blocks.append(("li",
        f"2) 典型负载性能（权重40%，T = (控制+感知+通信+运维)/4）："
        f"实时控制类 {_fmt(cats['realtime_control'], 2)} 分，"
        f"感知计算类 {_fmt(cats['perception'], 2)} 分，"
        f"通信协同类 {_fmt(cats['communication'], 2)} 分，"
        f"运维监控类 {_fmt(cats['monitoring'], 2)} 分，四类各占25%，"
        f"典型负载性能总得分 {_fmt(score['T'], 2)} 分。"))
    blocks.append(("li",
        f"3) 实时性能（权重30%，RT = 0.5×基础实时 + 0.5×调度）："
        f"基础实时性能 {_fmt(score['base_rt'], 2)} 分，"
        f"可调度性能 {_fmt(score['sched'], 2)} 分，两部分各占50%，"
        f"实时性能总得分 {_fmt(score['RT'], 2)} 分。"))
    blocks.append(("li", f"4) 功耗（权重10%，P = 四项工况均值）：综合得分 {_fmt(score['P'], 2)} 分。"))
    blocks.append(("p",
        f"最终按 0.2×{score['F']:.2f} + 0.3×{_fmt(score['RT'], 2)} + "
        f"0.4×{_fmt(score['T'], 2)} + 0.1×{_fmt(score['P'], 2)} 计算，"
        f"{board}的综合评分为 {_fmt(score['total'], 2)} 分。"))
    blocks.append(("p", "说明：本次评分对经过 min-max 归一化的分项采用 60 + 0.4×原分 的保底映射，"
                        "调度得分为 100×(1−平均截止时间错失率)，不参与该映射。"))
    return blocks


def blocks_to_markdown(blocks):
    out = []
    for kind, text in blocks:
        if kind == "h1":
            out.append("# " + text)
        elif kind == "h2":
            out.append("## " + text)
        elif kind == "h3":
            out.append("### " + text)
        else:
            out.append(text)
        out.append("")
    return "\n".join(out)


def write_docx(blocks, path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    for kind, text in blocks:
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "h3":
            doc.add_heading(text, level=3)
        else:
            para = doc.add_paragraph(text)
            if kind == "li":
                para.paragraph_format.left_indent = Pt(18)
    doc.save(path)


def safe_name(name):
    return re.sub(r"[^0-9A-Za-z一-鿿]+", "_", name)[:60]


def main():
    parser = argparse.ArgumentParser(description="生成 SSOT JSON 与自然语言描述")
    parser.add_argument("excel")
    parser.add_argument("--workloads", type=int, choices=[7, 9], default=9)
    parser.add_argument("-o", "--out-dir", default="reports")
    parser.add_argument("--scores-json", default=None)
    args = parser.parse_args()

    records = scoring.parse_blocks(args.excel)
    rows = scoring.compute(records, args.workloads)
    scores = {}
    for r in rows:
        scores[r["name"]] = {
            "base_rt": None if r["base_rt"] is None else scoring._floor(r["base_rt"]),
            "sched": r["sched"],
            "RT": r["RT"], "T": r["T"], "P": r["P"], "F": r["F"],
            "total": r["total"], "categories": r["categories"],
        }

    meta = json.load(open(META_PATH, encoding="utf-8")) if os.path.exists(META_PATH) else {}
    override = {}
    if args.scores_json:
        override = json.load(open(args.scores_json, encoding="utf-8"))
    boards = excel2json.parse_excel(args.excel, args.workloads)
    ssot_dir = os.path.join(args.out_dir, "ssot")
    text_dir = os.path.join(args.out_dir, "narrative")
    docx_dir = os.path.join(args.out_dir, "docx")
    for d in (ssot_dir, text_dir, docx_dir):
        os.makedirs(d, exist_ok=True)

    for board in boards:
        data = excel2json.to_result_json(board, args.excel, args.workloads)
        entry = meta.get(data["env"]["os_name"], {})
        data["env"]["os_version"] = entry.get("os_version") or ""
        families = entry.get("bsp_version") or {}
        family = next((k for k in families if k in data["env"]["board"]), None)
        data["env"]["bsp_version"] = families.get(family, "") if family else ""
        score = override.get(board["board"]) or scores.get(board["board"])
        if score is None:
            continue
        data["scores"] = {**data.get("scores", {}), **score}
        name = safe_name(board["board"])
        with open(os.path.join(ssot_dir, name + ".json"), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        blocks = build_narrative(data, score)
        with open(os.path.join(text_dir, name + ".md"), "w", encoding="utf-8") as f:
            f.write(blocks_to_markdown(blocks))
        write_docx(blocks, os.path.join(docx_dir, name + ".docx"))
    print(f"done: {len(boards)} boards -> {args.out_dir}")


if __name__ == "__main__":
    main()
