#!/usr/bin/env python3
"""把「结果展示表」Excel 逆生成结构化 result.json（SSOT）。

一个板卡块生成一份 JSON。缺失的字段写 null，并标注 source=excel_reverse。

用法：
  python3 utils/excel_to_result_json.py <结果展示表.xlsx> --workloads 9 -o out_dir
"""

import argparse
import json
import os
import re

import openpyxl

GENERIC = {"分类", "基", "性能", "功", "耗", "评", "典型", "负载", "实时"}
VENDOR_OS = {
    "翼辉": "SylixOS",
    "东土": "Intewell",
    "锐华": "ReWorks",
    "中移物联": "OneOS",
    "欧拉": "openEuler",
    "RTT": "RT-Thread",
}
OP_NAMES = {
    "信号量获取": "sem_take",
    "信号量释放": "sem_release",
    "消息发送": "mq_send",
    "消息接收": "mq_recv",
    "互斥锁获取": "mutex_take",
    "互斥锁释放": "mutex_release",
    "内存块申请": "mempool_alloc",
    "内存块释放": "mempool_free",
}
CATEGORY_NAMES = {
    "实时控制": "realtime_control",
    "感知计算": "perception",
    "通信协同": "communication",
    "运维监控": "monitoring",
}
ACCESS_OPS = ["rd", "wr", "cp", "frd", "fwr", "fcp", "memset", "memcpy"]


def _board_name(ws, start):
    for col in (7, 8, 1):
        value = ws.cell(start, col).value
        if not value:
            continue
        text = str(value).strip()
        if text and not text.startswith("表") and text not in GENERIC:
            return text
    return ""


def _num(value):
    if value is None:
        return None
    try:
        return float(str(value).replace("J", "").replace("ms", "").replace("us", "").strip())
    except ValueError:
        return None


def _parse_workloads(ws, start):
    out = []
    for offset, category in zip(range(9, 13), ["realtime_control", "perception", "communication", "monitoring"]):
        text = ws.cell(start + offset, 6).value
        if not text:
            continue
        for m in re.finditer(r"([A-Za-z][A-Za-z0-9]*)\s+([0-9.]+)\s*ms", str(text)):
            out.append({
                "name": m.group(1).lower(),
                "category": category,
                "exec_time_ms": float(m.group(2)),
                "avg_time_ms": None,
            })
    return out


def _parse_schedule(ws, sched_row, workload_count, cycles=3):
    gradients = {}
    text = None
    for r in range(sched_row, sched_row + 6):
        value = ws.cell(r, 9).value
        if isinstance(value, str) and value.strip().startswith("30%"):
            text = value
            break
    if text:
        for line in text.splitlines():
            m = re.match(r"\s*(\d+)%[：:]\s*([0-9.]+)", line.strip())
            if m:
                gradients[int(m.group(1))] = float(m.group(2))
    jobs = cycles * workload_count
    entries = []
    misses_total = 0
    for util in range(30, 101, 10):
        rate = gradients.get(util)
        misses = None if rate is None else round(rate * jobs)
        if misses is not None:
            misses_total += misses
        entries.append({
            "utilization_percent": util,
            "total_jobs": jobs,
            "deadline_misses": misses,
            "miss_rate": rate,
        })
    average = misses_total / (jobs * len(entries)) if entries else None
    return {
        "status": "from_excel",
        "config": {"cycles": cycles, "util_start": 30, "util_end": 100, "util_step": 10},
        "gradients": entries,
        "summary": {
            "average_miss_rate": average,
            "final_score": None if average is None else 100 * (1 - average),
        },
    }


def _parse_realtime(ws, start):
    service = []
    multicore = {"access": {}, "service": {}, "cross_core": {}}
    context_switch = None
    interrupt = None
    for r in range(start, start + 20):
        label = ws.cell(r, 11).value
        if label:
            label = str(label).strip()
            if label in OP_NAMES:
                service.append({
                    "operation": OP_NAMES[label],
                    "immediate_us": _num(ws.cell(r, 12).value),
                    "suspend_us": _num(ws.cell(r, 13).value),
                    "low_prio_us": _num(ws.cell(r, 14).value),
                    "high_prio_us": _num(ws.cell(r, 15).value),
                })
            elif label == "上下文切换延迟":
                context_switch = _num(ws.cell(r, 12).value)
            elif label == "中断软件延迟":
                interrupt = _num(ws.cell(r, 12).value)
        label17 = ws.cell(r, 17).value
        if label17:
            label17 = str(label17).strip()
            values = [ws.cell(r, c).value for c in (18, 19, 20, 21)]
            if label17 in ACCESS_OPS:
                multicore["access"][label17] = values
            elif label17 in ("任务间通信", "创建与删除"):
                multicore["service"][label17] = values
            elif label17 in ("同核通信", "异核通信"):
                multicore["cross_core"][label17] = values[0]
    return {
        "status": "from_excel",
        "single_core": {
            "context_switch": {"avg_us": context_switch},
            "interrupt": {"avg_us": interrupt},
            "syscall": {"avg_us": None},
            "service_cost": service,
        },
        "multicore": multicore,
    }


def _parse_power(ws, start):
    items = []
    for offset, name in zip(range(15, 19), ["static", "cpu", "memory", "file"]):
        value = ws.cell(start + offset, 6).value
        items.append({
            "item": name,
            "energy_j": _num(value),
            "score": ws.cell(start + offset, 5).value,
        })
    return items


def parse_excel(path, workload_count):
    wb = openpyxl.load_workbook(path, data_only=True)
    boards = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        vendor = sheet.split("-")[0].strip()
        for row in range(1, ws.max_row + 1):
            if str(ws.cell(row, 2).value or "").strip() != "调度性能":
                continue
            start = row - 14
            name = _board_name(ws, start)
            boards.append({
                "vendor": vendor,
                "os_name": VENDOR_OS.get(vendor, vendor),
                "board": name,
                "base_rt": ws.cell(start + 13, 5).value,
                "sched_score": ws.cell(row, 5).value,
                "realtime": _parse_realtime(ws, start),
                "schedule": _parse_schedule(ws, row, workload_count),
                "workloads": _parse_workloads(ws, start),
                "power": _parse_power(ws, start),
            })
    return boards


def to_result_json(board, source_file, workload_count):
    avg = board["schedule"]["summary"]["average_miss_rate"]
    return {
        "meta": {
            "framework_version": "1.0.0",
            "test_timestamp": None,
            "source": "excel_reverse",
            "source_file": os.path.basename(source_file),
            "workload_count": workload_count,
        },
        "env": {
            "os_name": board["os_name"],
            "os_version": None,
            "board": board["board"],
            "cpu_type": None,
            "cpu_freq_mhz": None,
            "cpu_core_num": None,
        },
        "modules": {
            "test-realtime": board["realtime"],
            "test-schedule": board["schedule"],
            "typical-workload": {"status": "from_excel", "workloads": board["workloads"]},
            "test-stress": {"status": "not_in_source"},
            "test-cmd": {"status": "not_in_source"},
        },
        "power": {"items": board["power"]},
        "scores": {
            "base_rt": board["base_rt"],
            "schedulability": board["sched_score"],
            "average_miss_rate": avg,
        },
    }


def main():
    parser = argparse.ArgumentParser(description="Excel 逆生成 result.json")
    parser.add_argument("excel")
    parser.add_argument("--workloads", type=int, choices=[7, 9], default=9)
    parser.add_argument("-o", "--out-dir", default="reports/ssot")
    args = parser.parse_args()

    boards = parse_excel(args.excel, args.workloads)
    os.makedirs(args.out_dir, exist_ok=True)
    index = []
    for board in boards:
        data = to_result_json(board, args.excel, args.workloads)
        safe = re.sub(r"[^0-9A-Za-z一-鿿]+", "_", board["board"])[:60]
        path = os.path.join(args.out_dir, f"{safe}.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        index.append({"board": board["board"], "os": board["os_name"], "file": path})
    with open(os.path.join(args.out_dir, "index.json"), "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)
    print(f"generated {len(index)} files in {args.out_dir}")


if __name__ == "__main__":
    main()
